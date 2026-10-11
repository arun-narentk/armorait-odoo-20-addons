# -*- coding: utf-8 -*-
import csv
import io
from collections import defaultdict

from odoo import Command, _, fields, models
from odoo.exceptions import UserError

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

COLUMNS = {
    'gstin': ('GSTIN of Deductor', 'GSTIN of E-commerce Operator', 'GSTIN of ECO', 'Deductor GSTIN', 'GSTIN'),
    'name': ("Deductor's Name", 'Deductor Name', 'E-commerce Operator Name', 'Trade/Legal Name', 'Name'),
    'taxable': ('Taxable Value', 'Net Amount Liable for TCS', 'Net Value', 'Amount Paid to Deductee', 'Value'),
    'igst': ('Integrated Tax', 'IGST'),
    'cgst': ('Central Tax', 'CGST'),
    'sgst': ('State/UT Tax', 'State Tax', 'SGST'),
}


class ArmoraGstTdsTcsImportWizard(models.TransientModel):
    _name = 'armora.gst.tds.tcs.import.wizard'
    _description = 'Import GSTR-2X TDS/TCS Credits'

    period_id = fields.Many2one('armora.gst.tds.tcs.period', required=True, ondelete='cascade')
    credit_type = fields.Selection([
        ('tds', 'GST TDS (section 51)'),
        ('tcs', 'GST TCS (section 52)'),
    ], string='Credits in the File', required=True, default='tds')
    data_file = fields.Binary(string='GSTR-2X File', required=True)
    filename = fields.Char()

    def action_import(self):
        self.ensure_one()
        period = self.period_id
        if period.state == 'closed':
            raise UserError(_('This month is closed.'))
        rows = self._armora_parse(self.data_file.content, self.filename or '')
        totals = defaultdict(lambda: defaultdict(float))
        names = {}
        for row in rows:
            gstin = (row.get('gstin') or '').strip().upper()
            if not gstin:
                continue
            names.setdefault(gstin, row.get('name') or '')
            for field in ('taxable', 'igst', 'cgst', 'sgst'):
                totals[gstin][field] += row.get(field) or 0.0
        if not totals:
            raise UserError(_('No deductor rows found. The file needs a GSTIN column and tax amounts.'))
        Partner = self.env['res.partner']
        for gstin, amounts in totals.items():
            credit = period.credit_ids.filtered(
                lambda c: c.credit_type == self.credit_type
                and ((c.deductor_gstin or '').upper() == gstin or (c.partner_id.vat or '').upper() == gstin)
            )[:1]
            vals = {
                'reported_taxable': amounts['taxable'],
                'reported_igst': amounts['igst'],
                'reported_cgst': amounts['cgst'],
                'reported_sgst': amounts['sgst'],
            }
            if credit:
                credit.write(vals)
                continue
            partner = Partner.search([
                ('vat', '=ilike', gstin), '|', ('company_id', '=', False), ('company_id', '=', period.company_id.id),
            ], limit=1)
            period.write({'credit_ids': [Command.create(dict(
                vals, credit_type=self.credit_type, deductor_gstin=gstin, partner_id=partner.commercial_partner_id.id or False,
                notes=False if partner else _('No contact with GSTIN %(gstin)s (%(name)s).', gstin=gstin, name=names[gstin]),
            ))]})
        period.message_post(body=_('GSTR-2X file %(file)s imported: %(count)s deductors.', file=self.filename or '', count=len(totals)))
        return {'type': 'ir.actions.act_window_close'}

    def _armora_parse(self, raw, filename):
        if filename.lower().endswith(('.xlsx', '.xls')):
            if load_workbook is None:
                raise UserError(_('Python package openpyxl is required to import Excel files.'))
            try:
                sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
            except Exception as err:
                raise UserError(_('Unable to open the Excel file: %s', err)) from err
            table = [list(row) for row in sheet.iter_rows(values_only=True)]
        else:
            try:
                text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = raw.decode('latin-1')
            table = list(csv.reader(io.StringIO(text)))
        header_index = next((i for i, row in enumerate(table[:20]) if any('gstin' in self._key(c) for c in row)), None)
        if header_index is None:
            raise UserError(_('No header row with a GSTIN column found.'))
        headers = [self._key(c) for c in table[header_index]]

        def column(aliases):
            for alias in aliases:
                if self._key(alias) in headers:
                    return headers.index(self._key(alias))
            for alias in aliases:
                for index, header in enumerate(headers):
                    if self._key(alias) in header:
                        return index
            return None

        mapping = {field: column(aliases) for field, aliases in COLUMNS.items()}
        rows = []
        for raw_row in table[header_index + 1:]:
            if not raw_row or all(v in (None, '') for v in raw_row):
                continue
            row = {}
            for field, index in mapping.items():
                value = raw_row[index] if index is not None and index < len(raw_row) else None
                row[field] = str(value).strip() if field in ('gstin', 'name') and value not in (None, '') else (
                    self._number(value) if field not in ('gstin', 'name') else False)
            rows.append(row)
        return rows

    @staticmethod
    def _key(value):
        return ''.join(ch for ch in str(value or '').lower() if ch.isalnum())

    @staticmethod
    def _number(value):
        if value in (None, ''):
            return 0.0
        try:
            return float(str(value).replace(',', '').strip())
        except ValueError:
            return 0.0
