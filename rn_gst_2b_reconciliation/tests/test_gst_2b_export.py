# -*- coding: utf-8 -*-
import io

from openpyxl import load_workbook

from odoo.tests import tagged
from odoo.tools.binary import BinaryBytes
from odoo.addons.l10n_in.tests.common import L10nInTestInvoicingCommon


@tagged('post_install', '-at_install')
class TestRnGst2bExport(L10nInTestInvoicingCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env.user.group_ids = [(4, cls.env.ref(
            'rn_gst_2b_reconciliation.group_rn_gst_2b_manager'
        ).id)]
        cls.default_company.vat = '24AAACB6992P1Z3'
        cls.vendor_gstin = '24AAPFU0939F1Z1'
        cls.partner_a.vat = cls.vendor_gstin
        cls.partner_a.l10n_in_gst_treatment = 'regular'

    def test_export_excel(self):
        bill = self.init_invoice(
            move_type='in_invoice', partner=self.partner_a, amounts=[1000.0],
            taxes=self.sgst_purchase_5, post=False, invoice_date='2025-04-15',
        )
        bill.ref = 'BILL-EXP-001'
        bill.invoice_line_ids.write({'l10n_in_hsn_code': '998314'})
        bill.action_post()
        recon = self.env['rn.gst.gstr2b'].create({
            'company_id': self.default_company.id,
            'gstin': self.default_company.vat,
            'financial_year': '2025-26',
            'return_period': '04',
            'period_start': '2025-04-01',
            'period_end': '2025-04-30',
        })
        csv = (
            'GSTIN,Invoice Number,Invoice Date,Taxable Value,CGST,SGST,IGST,Cess\n'
            '%(gstin)s,BILL-EXP-001,15-04-2025,1000,25,25,0,0\n'
            '%(gstin)s,ONLY-2B-777,10-04-2025,500,0,0,90,0\n'
        ) % {'gstin': self.vendor_gstin}
        self.env['rn.gst.2b.import.wizard'].create({
            'reconciliation_id': recon.id,
            'data_file': BinaryBytes(csv.encode('utf-8')),
            'filename': 'gstr2b.csv',
        }).action_import()
        recon.action_run_matching()

        action = recon.action_export_excel()
        self.assertEqual(action['type'], 'ir.actions.act_url')
        attachment = self.env['ir.attachment'].search([
            ('res_model', '=', 'rn.gst.gstr2b'), ('res_id', '=', recon.id),
        ], limit=1)
        self.assertTrue(attachment.file_size)

        book = load_workbook(io.BytesIO(attachment.raw))
        self.assertEqual(book.sheetnames, ['Summary', 'All Lines', 'Mismatch', 'Missing in Odoo', 'Missing in 2B'])
        self.assertEqual(book['All Lines'].max_row, 1 + len(recon.line_ids))
        missing = book['Missing in Odoo']
        self.assertEqual(missing.max_row, 2)
        self.assertEqual(missing.cell(row=2, column=5).value, 'ONLY-2B-777')
