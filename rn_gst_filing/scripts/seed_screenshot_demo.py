#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Seed armora20_gst_demo for Apps Store screenshots. Run via odoo shell."""
import io
import sys

from odoo.tools.binary import BinaryBytes

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

country_in = env.ref('base.in')
state_gj = env.ref('base.state_in_gj')
state_mh = env.ref('base.state_in_mh')
company = env.company

if company.account_fiscal_country_id.code != 'IN':
    company.write({
        'country_id': country_in.id,
    })
    env['account.chart.template'].try_loading('in', company, install_demo=True)

company.write({
    'name': 'Armora Technologies Pvt Ltd',
    'state_id': state_gj.id,
    'vat': '24AAACB6992P1Z3',
    'street': '100 Demo Tech Park',
    'city': 'Ahmedabad',
    'zip': '380001',
    'rn_gst_block_post_on_error': False,
    'rn_gst_recon_tolerance': 999999.0,
})

env.user.write({'company_id': company.id, 'company_ids': [(6, 0, [company.id])]})

Partner = env['res.partner']
partner_gj = Partner.search([('vat', '=', '24AAPFU0939F1Z1')], limit=1)
if not partner_gj:
    partner_gj = Partner.create({
        'name': 'Demo Customer Gujarat',
        'vat': '24AAPFU0939F1Z1',
        'state_id': state_gj.id,
        'country_id': country_in.id,
        'l10n_in_gst_treatment': 'regular',
    })
else:
    partner_gj.write({'l10n_in_gst_treatment': 'regular'})

partner_mh = Partner.search([('vat', '=', '27AAPFU0939F1ZV')], limit=1)
if not partner_mh:
    partner_mh = Partner.create({
        'name': 'Demo Customer Maharashtra',
        'vat': '27AAPFU0939F1ZV',
        'state_id': state_mh.id,
        'country_id': country_in.id,
        'l10n_in_gst_treatment': 'regular',
    })
else:
    partner_mh.write({'l10n_in_gst_treatment': 'regular'})

Product = env['product.product']
product = Product.search([('sale_ok', '=', True)], limit=1)
if product:
    product.write({'l10n_in_hsn_code': '998314'})

ChartTemplate = env['account.chart.template'].with_company(company)
sgst_18 = ChartTemplate.ref('sgst_sale_18', raise_if_not_found=False)
igst_18 = ChartTemplate.ref('igst_sale_18', raise_if_not_found=False)
if not sgst_18 or not igst_18:
    print('Indian sale taxes not found; ensure l10n_in chart is loaded.', file=sys.stderr)
    raise SystemExit(1)

Move = env['account.move']
existing = Move.search([
    ('company_id', '=', company.id),
    ('move_type', '=', 'out_invoice'),
    ('state', '=', 'posted'),
    ('invoice_date', '>=', '2025-04-01'),
    ('invoice_date', '<=', '2025-04-30'),
], limit=1)
if not existing:
    def _post_invoice(partner, taxes, amount, inv_date):
        move = Move.create({
            'move_type': 'out_invoice',
            'partner_id': partner.id,
            'invoice_date': inv_date,
            'l10n_in_state_id': partner.state_id.id or state_gj.id,
            'invoice_line_ids': [(0, 0, {
                'name': 'GST Demo Service',
                'product_id': product.id if product else False,
                'quantity': 1,
                'price_unit': amount,
                'tax_ids': [(6, 0, taxes.ids)],
                'l10n_in_hsn_code': '998314',
            })],
        })
        move.action_post()
        return move

    _post_invoice(partner_gj, sgst_18, 10000.0, '2025-04-10')
    _post_invoice(partner_mh, igst_18, 15000.0, '2025-04-15')
    _post_invoice(partner_gj, sgst_18, 5000.0, '2025-04-20')

Filing = env['rn.gst.filing']
filing = Filing.search([
    ('company_id', '=', company.id),
    ('financial_year', '=', '2025-26'),
    ('return_period', '=', '04'),
], limit=1)
if not filing:
    start, end = Filing.compute_period_dates('2025-26', '04', 'monthly')
    filing = Filing.create({
        'company_id': company.id,
        'gstin': company.vat,
        'financial_year': '2025-26',
        'return_period': '04',
        'filing_type': 'gstr1',
        'filing_frequency': 'monthly',
        'period_start': start,
        'period_end': end,
        'state_id': company.state_id.id,
    })

filing.action_load_data()
if filing.validation_status == 'pending':
    filing.action_validate()

Template = env['rn.gst.excel.template']
template = Template.search([('filing_type', '=', 'gstr1')], limit=1)
if not template and Workbook:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    sheets = {
        'b2b': ['GSTIN/UIN of Recipient', 'Receiver Name', 'Invoice Number', 'Invoice date', 'Invoice Value',
                'Place Of Supply', 'Reverse Charge', 'Invoice Type', 'Rate', 'Taxable Value',
                'Integrated Tax Paid', 'Central Tax Paid', 'State/UT Tax Paid', 'Cess Paid'],
        'b2cl': ['Invoice Number', 'Invoice date', 'Invoice Value', 'Place Of Supply', 'Rate', 'Taxable Value',
                 'Integrated Tax Paid', 'Cess Paid'],
        'b2cs': ['Type', 'Place Of Supply', 'Rate', 'Taxable Value', 'Integrated Tax Paid', 'Central Tax Paid',
                 'State/UT Tax Paid', 'Cess Paid'],
        'exp': ['Export Type', 'Invoice Number', 'Invoice date', 'Invoice Value', 'Port Code',
                'Shipping Bill Number', 'Shipping Bill Date', 'Rate', 'Taxable Value', 'Integrated Tax Paid'],
        'cdnr': ['GSTIN/UIN of Recipient', 'Receiver Name', 'Note Number', 'Note Date', 'Note Type',
                 'Place Of Supply', 'Reverse Charge', 'Note Value', 'Rate', 'Taxable Value',
                 'Integrated Tax Paid', 'Central Tax Paid', 'State/UT Tax Paid', 'Cess Paid'],
        'cdnur': ['UR Type', 'Note Number', 'Note Date', 'Note Type', 'Place Of Supply', 'Note Value',
                  'Rate', 'Taxable Value', 'Integrated Tax Paid', 'Cess Paid'],
        'hsn': ['HSN', 'Description', 'UQC', 'Total Quantity', 'Rate', 'Taxable Value',
                'Integrated Tax Amount', 'Central Tax Amount', 'State/UT Tax Amount', 'Cess Amount'],
        'nil': ['Description', 'Nil Rated Supplies', 'Exempted', 'Non-GST Supplies'],
        'docs': ['Nature of Document', 'Sr. No. From', 'Sr. No. To', 'Total Number', 'Cancelled'],
    }
    for name, headers in sheets.items():
        ws = wb.create_sheet(name)
        for col, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=header)
    buffer = io.BytesIO()
    wb.save(buffer)
    template = Template.create({
        'name': 'GSTR-1 Offline Tool Template (Demo)',
        'filing_type': 'gstr1',
        'version': '2025-26',
        'financial_year': '2025-26',
        'company_id': company.id,
        'template_file': BinaryBytes(buffer.getvalue()),
        'template_filename': 'gstr1_demo_template.xlsx',
        'active': True,
    })
    template.action_inspect_workbook()

if template:
    filing.template_id = template.id

if filing.validation_error_count == 0 and not filing.generated_file:
    try:
        filing.action_generate_excel()
    except Exception as exc:
        print('Excel generation note:', exc)

# Optional second filing for validation list screenshot (May period)
if not env['rn.gst.validation.result'].search_count([]):
    bad_filing = Filing.search([
        ('company_id', '=', company.id),
        ('return_period', '=', '05'),
    ], limit=1)
    if not bad_filing:
        start, end = Filing.compute_period_dates('2025-26', '05', 'monthly')
        bad_filing = Filing.create({
            'company_id': company.id,
            'gstin': company.vat,
            'financial_year': '2025-26',
            'return_period': '05',
            'filing_type': 'gstr1',
            'filing_frequency': 'monthly',
            'period_start': start,
            'period_end': end,
            'state_id': company.state_id.id,
        })
    if not bad_filing.line_ids:
        bad_filing.action_load_data()
    if bad_filing.line_ids:
        bad_filing.action_validate()

env.cr.commit()
print('READY filing_id=', filing.id, 'template_id=', template.id if template else None,
      'lines=', len(filing.line_ids), 'state=', filing.state,
      'errors=', filing.validation_error_count, 'generated=', bool(filing.generated_file))
