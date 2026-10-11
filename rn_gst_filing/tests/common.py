# -*- coding: utf-8 -*-
import io

from odoo.addons.l10n_in.tests.common import L10nInTestInvoicingCommon
from odoo.tools.binary import BinaryBytes

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None


class RnGstFilingTestCommon(L10nInTestInvoicingCommon):
    """Shared helpers: ACL groups, period dating, valid demo GSTINs."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        group_admin = cls.env.ref('rn_gst_filing.group_rn_gst_filing_administrator')
        group_base = cls.env.ref('rn_gst_base.group_rn_gst_administrator')
        cls.env.user.group_ids = [(4, group_admin.id), (4, group_base.id)]
        # Odoo l10n_in demo GSTIN may fail checksum; use known-valid GSTINs for filing tests.
        cls.valid_company_gstin = '24AAACB6992P1Z3'
        cls.valid_partner_gstin_gj = '24AAPFU0939F1Z1'
        cls.valid_partner_gstin_mh = '27AAPFU0939F1ZV'
        cls.default_company.vat = cls.valid_company_gstin
        cls.partner_a.vat = cls.valid_partner_gstin_gj
        cls.partner_b.vat = cls.valid_partner_gstin_mh
        cls.partner_a.l10n_in_gst_treatment = 'regular'
        cls.partner_b.l10n_in_gst_treatment = 'regular'
        if hasattr(cls, 'product_a'):
            cls.product_a.l10n_in_hsn_code = cls.product_a.l10n_in_hsn_code or '998314'
        # Filing tests validate at snapshot level; do not block posting via invoice validator.
        cls.default_company.write({
            'rn_gst_block_post_on_error': False,
            'rn_gst_recon_tolerance': 999999.0,
        })

    @classmethod
    def _create_gst_filing(cls, period='04', financial_year='2025-26', frequency='monthly'):
        company = cls.default_company
        start, end = cls.env['rn.gst.filing'].compute_period_dates(
            financial_year, period, frequency
        )
        registration = cls.env['rn.gst.registration'].search([
            ('company_id', '=', company.id),
            ('gstin', '=', company.vat),
        ], limit=1)
        return cls.env['rn.gst.filing'].create({
            'company_id': company.id,
            'gst_registration_id': registration.id if registration else False,
            'gstin': company.vat,
            'financial_year': financial_year,
            'return_period': period,
            'filing_type': 'gstr1',
            'filing_frequency': frequency,
            'period_start': start,
            'period_end': end,
            'state_id': company.state_id.id,
        })

    @classmethod
    def _make_template_xlsx(cls):
        if Workbook is None:
            return None
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        sheets = {
            'b2b': ['GSTIN/UIN of Recipient', 'Receiver Name', 'Invoice Number', 'Invoice date', 'Invoice Value',
                    'Place Of Supply', 'Reverse Charge', 'Invoice Type', 'Rate', 'Taxable Value',
                    'Integrated Tax Paid', 'Central Tax Paid', 'State/UT Tax Paid', 'Cess Paid'],
            'b2cl': ['Invoice Number', 'Invoice date', 'Invoice Value', 'Place Of Supply', 'Rate', 'Taxable Value',
                     'Integrated Tax Paid', 'Cess Paid'],
            'b2cs': ['Type', 'Place Of Supply', 'Rate', 'Taxable Value', 'Integrated Tax Paid', 'Central Tax Paid',
                     'State/UT Tax Paid', 'Cess Paid'],
            'hsnb2b': ['HSN', 'Description', 'UQC', 'Total Quantity', 'Rate', 'Taxable Value',
                       'Integrated Tax Amount', 'Central Tax Amount', 'State/UT Tax Amount', 'Cess Amount'],
            'docs': ['Nature of Document', 'Sr. No. From', 'Sr. No. To', 'Total Number', 'Cancelled'],
        }
        for name, headers in sheets.items():
            ws = wb.create_sheet(name)
            for col, header in enumerate(headers, start=1):
                ws.cell(row=1, column=col, value=header)
        buffer = io.BytesIO()
        wb.save(buffer)
        return BinaryBytes(buffer.getvalue())
