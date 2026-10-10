# -*- coding: utf-8 -*-
import io
import json

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


class ArmoraGstExcelTemplate(models.Model):
    _name = 'armora.gst.excel.template'
    _description = 'GSTN Returns Offline Tool Excel Template'
    _order = 'filing_type, financial_year desc, version desc, id desc'

    name = fields.Char(required=True)
    filing_type = fields.Selection([
        ('gstr1', 'GSTR-1'),
        ('gstr1a', 'GSTR-1A'),
        ('gstr3b', 'GSTR-3B'),
        ('gstr9', 'GSTR-9'),
    ], required=True, default='gstr1', index=True)
    version = fields.Char(
        required=True,
        default='1',
        help='Label matching the Offline Tool / template release you use (GSTN may change templates).',
    )
    financial_year = fields.Char(
        string='Financial Year',
        help='Optional FY label such as 2025-26',
    )
    active = fields.Boolean(default=True)
    effective_from = fields.Date()
    effective_to = fields.Date()
    template_file = fields.Binary(
        string='Offline Tool Workbook',
        required=True,
        attachment=True,
        help='Official Excel template from the GSTN Returns Offline Tool (GSTR-1/IFF). '
             'Administrators should re-upload when GSTN publishes a new template version.',
    )
    template_filename = fields.Char()
    mapping_json = fields.Text(
        string='Sheet/Column Mapping JSON',
        help='Maps logical GSTR sections to sheet names and header columns discovered in the '
             'uploaded Offline Tool workbook.',
    )
    notes = fields.Text(
        help='Optional notes, e.g. Offline Tool build date or GSTN template release.',
    )
    company_id = fields.Many2one(
        'res.company',
        string='Company',
        default=lambda self: self.env.company,
        index=True,
    )
    sheet_names = fields.Char(string='Detected Sheets', readonly=True)
    validation_message = fields.Text(readonly=True)

    _armora_gst_template_uniq = models.Constraint(
        'unique(company_id, filing_type, version, financial_year)',
        "A template with the same filing type, FY and version already exists for this company.",
    )

    @api.model
    def _default_section_keys(self):
        return [
            'b2b', 'b2cl', 'b2cs', 'exp', 'cdnr', 'cdnur',
            'hsn_b2b', 'hsn_b2c', 'nil', 'docs',
        ]

    def action_inspect_workbook(self):
        self.ensure_one()
        if not self.template_file:
            raise UserError(_(
                'Upload the official GSTN Returns Offline Tool Excel workbook first.'
            ))
        if load_workbook is None:
            raise UserError(_('Python package openpyxl is required to inspect Excel templates.'))
        try:
            data = self.template_file.content
            wb = load_workbook(filename=io.BytesIO(data), data_only=False, read_only=True)
        except Exception as err:
            raise UserError(_('Unable to open Offline Tool workbook: %s') % err) from err

        sheets = list(wb.sheetnames)
        if not sheets:
            raise ValidationError(_('Offline Tool workbook has no sheets.'))
        if len(sheets) != len(set(sheets)):
            raise ValidationError(_('Offline Tool workbook has duplicate sheet names.'))

        mapping = {}
        issues = []
        for sheet_name in sheets:
            ws = wb[sheet_name]
            headers = []
            # Scan first 20 rows for a header-like row (non-empty cells)
            header_row_idx = None
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
                values = [str(c).strip() if c is not None else '' for c in row]
                non_empty = [v for v in values if v]
                if len(non_empty) >= 2:
                    headers = values
                    header_row_idx = row_idx
                    break
            if not headers:
                issues.append(_('Sheet "%s": no header row detected in first 20 rows.') % sheet_name)
                continue
            # duplicate headers
            cleaned = [h for h in headers if h]
            if len(cleaned) != len(set(cleaned)):
                issues.append(_('Sheet "%s": duplicate header labels found.') % sheet_name)
            mapping[sheet_name] = {
                'header_row': header_row_idx,
                'headers': headers,
                'logical_section': self._guess_logical_section(sheet_name, headers),
            }
        wb.close()

        # Ensure each logical section can be mapped if sheet names look official
        logical = {v.get('logical_section') for v in mapping.values() if v.get('logical_section')}
        missing = [s for s in self._default_section_keys() if s not in logical]
        if 'hsn_b2b' in missing and 'hsn_b2c' in missing:
            if 'hsn_b2b' in logical or 'hsn_b2c' in logical:
                missing = [s for s in missing if s not in ('hsn_b2b', 'hsn_b2c')]
            else:
                missing = [s for s in missing if not s.startswith('hsn_')]
                missing.append('hsn (B2B or B2C)')
        if missing:
            issues.append(
                _(
                    'Could not auto-map logical sections: %s. '
                    'Update mapping_json after upload to match the Offline Tool template.'
                )
                % ', '.join(missing)
            )

        self.write({
            'sheet_names': ', '.join(sheets),
            'mapping_json': json.dumps(mapping, indent=2),
            'validation_message': (
                '\n'.join(issues) if issues
                else _('Offline Tool workbook inspected successfully.')
            ),
        })
        return True

    @api.model
    def _guess_logical_section(self, sheet_name, headers):
        """Map Offline Tool sheet names to logical GSTR-1 sections.

        Prefer exact GSTN sheet names. When matching aliases, check longer
        keys first so ``cdnur`` does not match ``cdnr`` and ``b2cl`` does not
        collapse into ``b2b``.
        """
        key = (sheet_name or '').strip().lower().replace(' ', '')
        # Common official GSTR-1 / IFF Offline Tool sheet names (exact).
        exact = {
            'b2b': 'b2b',
            'b2cl': 'b2cl',
            'b2cs': 'b2cs',
            'exp': 'exp',
            'cdnr': 'cdnr',
            'cdnur': 'cdnur',
            'hsnb2b': 'hsn_b2b',
            'hsnb2c': 'hsn_b2c',
            'hsn': 'hsn_b2b',
            'exemp': 'nil',
            'nil': 'nil',
            'docs': 'docs',
            'documents': 'docs',
        }
        if key in exact:
            return exact[key]

        # Longer aliases first to avoid substring collisions.
        aliases = [
            ('cdnur', ['cdnur']),
            ('cdnr', ['cdnr']),
            ('b2cl', ['b2cl']),
            ('b2cs', ['b2cs']),
            ('b2b', ['b2b']),
            ('exp', ['exp', 'export']),
            ('hsn_b2b', ['hsn(b2b)', 'hsnb2b']),
            ('hsn_b2c', ['hsn(b2c)', 'hsnb2c']),
            ('hsn_b2b', ['hsn']),
            ('nil', ['exemp', 'nill', 'nil']),
            ('docs', ['documents', 'docs', 'doc']),
        ]
        for section, names in aliases:
            if any(key == a or key.startswith(a) for a in names):
                # Skip amendment sheets (b2ba, cdnra, …) for V1 mapping.
                if key.endswith('a') and key != 'exemp' and len(key) > 1 and key[:-1] in exact:
                    return False
                return section

        header_blob = ' '.join(h.lower() for h in headers if h)
        if 'gstin' in header_blob and 'invoice' in header_blob:
            return 'b2b'
        return False

    @api.model
    def get_active_template(self, company, filing_type='gstr1', on_date=None):
        domain = [
            ('filing_type', '=', filing_type),
            ('active', '=', True),
            ('template_file', '!=', False),
            '|',
            ('company_id', '=', False),
            ('company_id', '=', company.id),
        ]
        templates = self.search(domain, order='company_id desc, effective_from desc, id desc')
        on_date = on_date or fields.Date.context_today(self)
        for tmpl in templates:
            if tmpl.effective_from and tmpl.effective_from > on_date:
                continue
            if tmpl.effective_to and tmpl.effective_to < on_date:
                continue
            return tmpl
        return self.browse()
